# -*- coding: utf-8 -*-
"""把导师录音转写文本 teacher_transcript.txt 转成 Word 文档。"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / ".tmp_audio" / "teacher_transcript.txt"
OUT = ROOT / "doc" / "导师录音转写_2026-10-03.docx"

doc = Document()

# 标题
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("导师录音转写文本")
r.font.size = Pt(18)
r.font.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("SEU 旋转机械故障诊断项目 · 2026-10-03 录音（faster-whisper medium 自动转写）")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

note = doc.add_paragraph()
r = note.add_run("说明：以下为自动语音识别（ASR）转写结果，保留原始时间戳，可能存在少量错别字与同音字误差（如“造势”=噪声、“性造比”=信噪比、“试一图”=示意图等），仅供内部参考。")
r.font.size = Pt(9)
r.font.italic = True
r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

lines = SRC.read_text(encoding="utf-8").strip().splitlines()
for line in lines:
    line = line.strip()
    if not line:
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(line)
    run.font.size = Pt(10.5)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(f"已生成: {OUT}")
print(f"共 {len(lines)} 行")
