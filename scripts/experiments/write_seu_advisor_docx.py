"""Advisor Word: 思路, 方法, 参数, 结果, 分析. Technical details, not source."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "doc" / "handoff" / "seu_stage_report.docx"
FIG = ROOT / "doc" / "figures" / "seu"
EAST_ASIA = "宋体"
LATIN = "Times New Roman"


def set_run_font(run, size=12, bold=False) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA)


def add_para(doc, text, *, size=12, bold=False, space_after=8, indent=True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_heading_cn(doc, text, level=1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)


def add_table(doc, header, rows) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for i, name in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(name))
        set_run_font(run, size=10, bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def add_figure(doc, filename: str, caption: str) -> None:
    path = FIG / filename
    if path.exists():
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.space_after = Pt(4)
        pic.paragraph_format.first_line_indent = Cm(0)
        pic.add_run().add_picture(str(path), width=Inches(5.5))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = cap.add_run(caption)
    set_run_font(run, size=10)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA)


def build() -> Path:
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("多时间步状态转移网络与 CNN / GCN 故障分类")
    set_run_font(r, size=18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)
    r2 = sub.add_run("东南大学 DDS 齿轮箱")
    set_run_font(r2, size=14, bold=True)

    add_heading_cn(doc, "1  思路")
    add_para(
        doc,
        "复杂网络方法把振动变成状态之间的转移，再用节点度、路径长度、方差一类指标去做分类。"
        "指标靠人来选，选漏了就没有补救；只记相邻一步转移，更长间隔上的动态进不了模型；"
        "论文里的网络图往往是示意，真正计算用的是矩阵。",
    )
    add_para(
        doc,
        "本工作把人工选指标换成模型自动提特征。振动按幅值离散成有序状态，状态之间的跳转做成有向加权矩阵；"
        "一步、两步、三步各一张，叠在一起，分别交给 CNN 和 GCN。"
        "CNN 把矩阵当多通道小图卷积，GCN 把状态当节点、把转移当边。两条骨干互斥。"
        "注意力只作为开关：CNN 是卷积前的空间自注意力，GCN 是边上的图注意力，二者不是同一模块。",
    )
    add_para(
        doc,
        "节点按幅值从低到高编号，分档边界只在该任务、该工况的训练集上确定，验证与测试沿用同一把尺子。"
        "这样不同窗里的同一编号表示同一段幅值，矩阵的行列才有固定含义。",
    )
    add_para(
        doc,
        "实验对象是东南大学 DDS 齿轮箱。轴承与齿轮各建一套五分类，两种工况分开训练，再对两个验证准确率取平均。"
        "振动只用行星齿轮箱三轴。测试集按时间切出，本阶段不评分，避免用测试去选结构或超参。",
    )

    add_heading_cn(doc, "2  方法")
    add_para(
        doc,
        "轴承和齿轮各做一套五分类，标签不混用。"
        "工况 20-0 与 30-2 各自切分、各自确定幅值分档、各自训练，再对两个验证准确率取算术平均。"
        "振动先做成转移矩阵，CNN 与 GCN 读同一套矩阵，一次训练只走其中一种网络。",
    )
    add_para(
        doc,
        "处理链是：原始记录 → 行星齿轮箱三轴振动 → 按时间切开的训练、验证、测试段 → 不重叠窗 → "
        "用训练集定幅值分档 → 符号序列 → 多时间步有向转移矩阵 → CNN 或 GCN → 五类。"
        "取值见第 3 节。",
    )

    add_heading_cn(doc, "2.1 原始记录变成三轴振动", level=2)
    add_para(
        doc,
        "数据是东南大学 DDS 齿轮箱。轴承五类：健康、滚动体、内圈、外圈、复合。"
        "齿轮五类：健康、缺齿、丢齿、齿根、表面。"
        "工况记为 20-0 和 30-2。",
    )
    add_para(
        doc,
        "每条记录 1 048 560 个采样点。原始表 8 列，只用行星齿轮箱 x、y、z 三轴，变成长度 1 048 560 的三条振动。"
        "转矩和其他振动不用。不滤波，不按通道标准化。标称采样率 5120 Hz，本阶段按点数切窗，不按秒。",
    )

    add_heading_cn(doc, "2.2 先按时间切开，再切成窗", level=2)
    add_para(
        doc,
        "一条记录按时间切成前 70% 训练、接着 15% 验证、最后 15% 测试。三段相连，互不交叉。"
        "切分按时间比例，不随机打乱，也不在窗之间交叉验证。",
    )
    add_para(
        doc,
        "再在各段内切窗。窗长 800 点，步长 800 点，不重叠，窗不跨切点。每个窗是 3 轴 × 800 点。"
        "每个文件、每一类得到训练 917 窗、验证 196 窗、测试 196 窗。"
        "五类合计：训练 4585，验证 980，测试 980。测试窗已切出，本阶段不评分。",
    )

    add_heading_cn(doc, "2.3 窗变成符号序列", level=2)
    add_para(
        doc,
        "每个通道把幅值等宽分成 6 档，得到 6 个有序状态。"
        "分档边界只用该任务、该工况的全部训练窗：按通道取最小、最大值，中间均分成 6 段。"
        "验证和测试不准重算边界。两个工况各自定边界，不共用。",
    )
    add_para(
        doc,
        "一个窗里，每个通道的 800 个点落入 0 到 5，变成一条长度 800 的整数序列。"
        "编号按幅值从低到高：0 是最低档，5 是最高档。矩阵第 i 行第 j 列始终表示从第 i 档转到第 j 档。",
    )

    add_heading_cn(doc, "2.4 符号序列变成转移矩阵", level=2)
    add_para(
        doc,
        "对一条符号序列，分别统计滞后 1、2、3 个点的有向转移，各得一张 6×6 矩阵。"
        "滞后 1 看相邻两点，滞后 2 看隔一个点，滞后 3 同理。格子先计次数，自己转到自己也计。"
        "再把整张矩阵的次数和除成 1，不是按行归一化。每个通道得到 3 张 6×6。",
    )
    add_para(
        doc,
        "三轴的三张图叠在一起，一个窗变成 9×6×6。"
        "顺序是 x 的 1、2、3 步，再 y 的 1、2、3 步，再 z 的 1、2、3 步。"
        "CNN 和 GCN 用同一份矩阵，不各算一套。",
    )

    add_heading_cn(doc, "2.5 矩阵变成类别", level=2)
    add_para(
        doc,
        "输入都是这份 9×6×6。一次只训 CNN 或只训 GCN。五类交叉熵：网络输出未归一化的类别分数，损失里做 Softmax。"
        "优化器 Adam。第一版学习率 5×10⁻⁴，权重衰减 10⁻⁴，dropout 0.2，批大小 32。"
        "验证准确率升高才保存参数；连续 30 轮不升则停；最多 200 轮。"
        "训练随机种子 42，只管初始化和训练打乱，不管时间切分。"
        "工况 20-0 训完存一份，30-2 另训另存。",
    )
    add_para(
        doc,
        "CNN：两层卷积，通道 32→64，核 3，池化 2，批归一化，ReLU。"
        "6×6 经过两次池化后空间收到约 1×1，再全局平均，接 Dropout 和线性层到 5 类。卷积后不再加隐层。"
        "CNN 的注意力若打开，是卷积之前的空间自注意力。第一版关闭。",
    )
    add_para(
        doc,
        "GCN：6 个幅值档是节点，6×6 是有向边权。"
        "9 张图拆成 3 个振动通道 × 3 个时间步，同一通道上三个步长作为三种关系一起传消息。"
        "两层，隐层 32。GCN 的注意力若打开，边权同时看两端节点特征和转移概率；关闭时按转移概率加权。"
        "第一版关闭。CNN 自注意力和 GCN 图注意力不是同一个模块。",
    )

    add_heading_cn(doc, "2.6 对照和扫参", level=2)
    add_para(
        doc,
        "四条线：CNN 无注意力、CNN 自注意力、GCN 无图注意力、GCN 图注意力。"
        "对照只改网络，不改已经做好的矩阵和切分。"
        "各线只扫学习率 {1×10⁻⁴, 5×10⁻⁴, 1×10⁻³} 和 dropout {0.1, 0.2, 0.3}。"
        "轴承、齿轮各四条线，每条线两个工况，共 144 组。"
        "选参看该任务两个工况验证准确率的平均。测试集全程不打开。",
    )

    add_heading_cn(doc, "2.7 图怎么来", level=2)
    add_para(
        doc,
        "混淆矩阵：验证集预测。行是真实类别，列是预测类别，每类 196 窗。"
        "t-SNE：分类头之前的向量降到二维，颜色是真实类别。"
        "CNN 用卷积之后、线性分类之前的全局向量；GCN 用图卷积之后、线性头之前的向量。"
        "下列图对应第一版，注意力关闭。打开注意力后的结果未重画。",
    )

    add_heading_cn(doc, "3  参数")
    add_heading_cn(doc, "3.1 这一版定下来的", level=2)
    add_para(doc, "下面这些在第一版定死，扫参时没有改。")
    add_table(
        doc,
        ["项", "取值"],
        [
            ["数据", "东南大学 DDS 齿轮箱"],
            ["任务", "轴承五分类、齿轮五分类，分开做"],
            ["轴承类别", "健康、滚动体、内圈、外圈、复合"],
            ["齿轮类别", "健康、缺齿、丢齿、齿根、表面"],
            ["工况", "20-0、30-2，分训后再对验证准确率取平均"],
            ["通道", "行星齿轮箱 x、y、z；不用转矩"],
            ["单条记录", "1 048 560 点；标称 5120 Hz"],
            ["预处理", "不滤波、不标准化"],
            ["划分", "时间 70% / 15% / 15%，不随机"],
            ["窗长 / 步长", "800 / 800，不重叠"],
            ["每类窗数", "训练 917，验证 196，测试 196"],
            ["一次训练样本", "训练 4585，验证 980，测试 980"],
            ["表示", "多时间步状态转移矩阵"],
            ["状态数", "6，等宽，边界只由该工况训练集确定"],
            ["时间步", "1、2、3"],
            ["边", "有向；自环保留；权重为全矩阵次数归一化"],
            ["多通道", "三轴堆叠，输入 9×6×6"],
            ["CNN 结构", "卷积通道 32→64，核 3，池化 2，无额外全连接隐层"],
            ["GCN 结构", "关系式布局，2 层，隐层 32，三步一起传消息"],
            ["优化器", "Adam，权重衰减 10⁻⁴，批大小 32"],
            ["早停", "按验证准确率，耐心 30，最多 200 轮"],
            ["训练种子", "42"],
            ["测试集", "已切出，未评分"],
        ],
    )
    add_heading_cn(doc, "3.2 扫过的", level=2)
    add_para(
        doc,
        "只扫学习率和 dropout。第一版是学习率 5×10⁻⁴、dropout 0.2、注意力关闭。"
        "四条线各自在学习率 {1×10⁻⁴, 5×10⁻⁴, 1×10⁻³} × dropout {0.1, 0.2, 0.3} 上选验证平均最好的一组。"
        "窗长、状态数、时间步、是否加权、是否有向、人工网络指标，都没有扫。",
    )

    add_heading_cn(doc, "4  结果")
    add_para(doc, "全部为验证准确率。测试集未评。平均是 20-0 与 30-2 的算术平均。")
    add_para(doc, "第一版（学习率 5×10⁻⁴，dropout 0.2，注意力关闭）：")
    add_table(
        doc,
        ["任务", "方法", "20-0", "30-2", "平均"],
        [
            ["轴承", "CNN", "0.6857", "0.9776", "0.8316"],
            ["轴承", "GCN", "0.7296", "0.9898", "0.8597"],
            ["齿轮", "CNN", "0.9194", "0.8582", "0.8888"],
            ["齿轮", "GCN", "0.9184", "0.8194", "0.8689"],
        ],
    )
    add_para(doc, "四条线各自选参后的最好一组：")
    add_table(
        doc,
        ["任务", "方法", "学习率", "dropout", "20-0", "30-2", "平均"],
        [
            ["轴承", "CNN", "5×10⁻⁴", "0.1", "0.6867", "0.9857", "0.8362"],
            ["轴承", "CNN + 自注意力", "1×10⁻³", "0.2", "0.6939", "0.9878", "0.8408"],
            ["轴承", "GCN", "5×10⁻⁴", "0.3", "0.7316", "0.9898", "0.8607"],
            ["轴承", "GCN + 图注意力", "5×10⁻⁴", "0.2", "0.7296", "0.9908", "0.8602"],
            ["齿轮", "CNN", "5×10⁻⁴", "0.1", "0.9286", "0.8622", "0.8954"],
            ["齿轮", "CNN + 自注意力", "1×10⁻³", "0.1", "0.9265", "0.8602", "0.8934"],
            ["齿轮", "GCN", "1×10⁻³", "0.3", "0.9184", "0.8541", "0.8862"],
            ["齿轮", "GCN + 图注意力", "1×10⁻³", "0.3", "0.9204", "0.8480", "0.8842"],
        ],
    )

    add_figure(
        doc,
        "bearing_cnn_20_0_confusion.png",
        "图 1　轴承 20-0，CNN，验证混淆矩阵。准确率 0.6857。滚动体、复合、健康交叉；内圈对角占优。",
    )
    add_figure(
        doc,
        "bearing_cnn_20_0_tsne.png",
        "图 2　轴承 20-0，CNN，分类头前 t-SNE。内圈成团；健康、滚动体、复合重叠。",
    )
    add_figure(
        doc,
        "bearing_cnn_30_2_confusion.png",
        "图 3　轴承 30-2，CNN，验证混淆矩阵。准确率 0.9776。对角占优。",
    )
    add_figure(
        doc,
        "bearing_cnn_30_2_tsne.png",
        "图 4　轴承 30-2，CNN，t-SNE。五类分离。与图 2 对比，差距来自工况。",
    )
    add_figure(
        doc,
        "bearing_gcn_20_0_confusion.png",
        "图 5　轴承 20-0，GCN，验证混淆矩阵。准确率 0.7296。与 CNN 同一套转移矩阵。内圈清楚，其余三类仍混。",
    )
    add_figure(
        doc,
        "bearing_gcn_20_0_tsne.png",
        "图 6　轴承 20-0，GCN，图卷积后 t-SNE。与图 5 同一现象。",
    )
    add_figure(
        doc,
        "bearing_gcn_30_2_confusion.png",
        "图 7　轴承 30-2，GCN，验证混淆矩阵。准确率 0.9898。复合类 196 窗全部正确。",
    )
    add_figure(
        doc,
        "bearing_gcn_30_2_tsne.png",
        "图 8　轴承 30-2，GCN，t-SNE。五类分离。",
    )
    add_figure(
        doc,
        "gear_cnn_20_0_confusion.png",
        "图 9　齿轮 20-0，CNN，验证混淆矩阵。准确率 0.9194。健康、缺齿几乎全对；误差在缺齿与表面。",
    )
    add_figure(
        doc,
        "gear_cnn_20_0_tsne.png",
        "图 10　齿轮 20-0，CNN，t-SNE。健康远离其余类。",
    )
    add_figure(
        doc,
        "gear_cnn_30_2_confusion.png",
        "图 11　齿轮 30-2，CNN，验证混淆矩阵。准确率 0.8582。健康全对；缺齿、齿根、表面互串。",
    )
    add_figure(
        doc,
        "gear_cnn_30_2_tsne.png",
        "图 12　齿轮 30-2，CNN，t-SNE。健康独立，其余四类靠近。",
    )
    add_figure(
        doc,
        "gear_gcn_20_0_confusion.png",
        "图 13　齿轮 20-0，GCN，验证混淆矩阵。准确率 0.9184。与 CNN 接近。",
    )
    add_figure(
        doc,
        "gear_gcn_20_0_tsne.png",
        "图 14　齿轮 20-0，GCN，t-SNE。健康分离，故障类有重叠。",
    )
    add_figure(
        doc,
        "gear_gcn_30_2_confusion.png",
        "图 15　齿轮 30-2，GCN，验证混淆矩阵。准确率 0.8194。第一版齿轮 30-2 最低。",
    )
    add_figure(
        doc,
        "gear_gcn_30_2_tsne.png",
        "图 16　齿轮 30-2，GCN，t-SNE。健康一侧，故障类叠在另一侧。",
    )

    add_heading_cn(doc, "5  分析")
    add_para(
        doc,
        "当前停在第一版表示上：窗长 800、6 个状态、时间步 1/2/3、有向、概率权重、三通道堆叠。"
        "已经试过的是：换 CNN 或 GCN、开或关注意力、扫学习率和 dropout。矩阵怎么从振动算出来，这一版没有改。",
    )
    add_para(
        doc,
        "第一，瓶颈在表示与工况，不在学习率。144 组只把平均抬了约 0.001–0.017。"
        "轴承 20-0 仍停在 0.69–0.73，30-2 已到 0.98–0.99。各类窗数相同，不是类别不平衡。"
        "同一套矩阵换工况就可以分开，说明 20-0 上健康、滚动体、复合这三类在转移矩阵里本身叠在一起。"
        "图 2 与图 4、图 6 与图 8 对看：分类头之前就已经分不开，再扫 dropout 补不回来。",
    )
    add_para(
        doc,
        "第二，注意力这一版没有拉开差距。CNN 的自注意力加在 9×6×6 的小空间上；GCN 的边本来已带转移概率。"
        "扫参后开注意力的平均与关闭时相差不到 0.01，齿轮上甚至略低。"
        "按「有效再写进创新、无效不必强留」，目前不能把注意力写成主要贡献。",
    )
    add_para(
        doc,
        "第三，CNN 与 GCN 没有压倒性优劣。轴承上 GCN 平均高约 0.02，主要来自 20-0（0.73 对 0.69）；"
        "齿轮上 CNN 平均更高，30-2 上 GCN 更弱（0.82–0.85 对 CNN 的 0.86）。"
        "两者读同一份矩阵，差的是读法，不是两套数据。",
    )
    add_para(
        doc,
        "第四，错误有结构。轴承 20-0：内圈可分，难在健康 / 滚动体 / 复合。"
        "齿轮：健康几乎总是单独成团；难在缺齿、齿根、表面，困难工况是 30-2，与轴承相反。"
        "后面若改窗长、状态数或时间步，应看这些易混类是否分开，而不是只看总平均。",
    )
    add_para(
        doc,
        "第五，测试集未评，上面不能当成测试性能。"
        "人工网络指标对照、只做一步转移、改节点数与窗长都还没做。",
    )

    add_heading_cn(doc, "6  引用")
    add_para(
        doc,
        "数据：东南大学 DDS 齿轮箱，轴承与齿轮两组记录，工况 20-0、30-2。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(OUT))
        return OUT
    except PermissionError:
        alt = OUT.with_name("seu_stage_report_updated.docx")
        doc.save(str(alt))
        return alt


if __name__ == "__main__":
    path = build()
    print(path, path.stat().st_size)
