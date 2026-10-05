"""导出「本文 99.85%~100%」对应的轴承实验原始记录 -> xlsx（openpyxl 直写，数值为实测字面值）。"""
from __future__ import annotations

import csv
import json
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

SEU = Path(r"C:/Users/lgt11/Desktop/EEG/EEG/datas/experiments_seu")
OUT = Path(r"C:/Users/lgt11/Desktop/EEG/EEG/doc/references/轴承实验数据_本文99.85-100.xlsx")
TITLE = "轴承实验数据导出｜本文 99.85% ~ 100%（SEU DDS 轴承五分类）"

CLASSES = ["health", "ball", "inner", "outer", "comb"]
CLASS_CN = {"health": "正常", "ball": "滚动体", "inner": "内圈", "outer": "外圈", "comb": "内外圈复合"}
COND_CN = {"20_0": "20 Hz-0 V", "30_2": "30 Hz-2 V"}
COMBOS = [("cnn", "20_0"), ("cnn", "30_2"), ("gcn", "20_0"), ("gcn", "30_2")]
ARCH_CN = {"cnn": "CNN", "gcn": "GCN"}


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEAD = xl_color("#4472C4")
XL_ACCENT = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
XL_OK_BG = xl_color("#C6EFCE")
XL_OK_FG = xl_color("#006100")
XL_WARN_BG = xl_color("#FFEB9C")
XL_WARN_FG = xl_color("#9C6500")
XL_BAD_BG = xl_color("#FFC7CE")
XL_BAD_FG = xl_color("#9C0006")
XL_TITLE_FG = xl_color("#1F3864")

F_TITLE = Font(name="微软雅黑", size=13, bold=True, color=XL_TITLE_FG)
F_SUB = Font(name="微软雅黑", size=10, color=xl_color("#404040"))
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color=xl_color("#FFFFFF"))
F_BODY = Font(name="微软雅黑", size=10)
F_BOLD = Font(name="微软雅黑", size=10, bold=True)
F_BLOCK = Font(name="微软雅黑", size=11, bold=True, color=XL_TITLE_FG)

FILL_HEAD = PatternFill("solid", fgColor=XL_HEAD)
FILL_ACCENT = PatternFill("solid", fgColor=XL_ACCENT)
FILL_TOTAL = PatternFill("solid", fgColor=xl_color("#F2F2F2"))
FILL_DIAG = PatternFill("solid", fgColor=xl_color("#EDF3FB"))

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

AL_C = Alignment(horizontal="center", vertical="center")
AL_L = Alignment(horizontal="left", vertical="center")
AL_LW = Alignment(horizontal="left", vertical="top", wrap_text=True)


def load_run(arch: str, cond: str) -> dict:
    base = SEU / f"checkpoints_bearing_{arch}_resplit" / "bearing" / cond / "seed_42"
    m = json.loads((base / "metrics.json").read_text(encoding="utf-8"))
    preds = list(csv.DictReader((base / "test_predictions.csv").open(encoding="utf-8")))
    cm = [[0] * 5 for _ in range(5)]
    for r in preds:
        cm[int(r["true_label"])][int(r["pred_label"])] += 1
    assert cm == m["test_confusion_matrix"], f"confusion matrix mismatch: {arch} {cond}"
    n = len(preds)
    correct = sum(1 for r in preds if r["true_label"] == r["pred_label"])
    assert correct / n == m["test_acc"], f"acc mismatch: {arch} {cond}"
    return {"arch": arch, "cond": cond, "m": m, "cm": cm, "preds": preds, "n": n, "correct": correct}


def load_cv(arch: str, cond: str) -> dict:
    return json.loads((SEU / "cv" / f"bearing_{arch}_{cond}.json").read_text(encoding="utf-8"))


RUNS = [load_run(a, c) for a, c in COMBOS]
CVS = [load_cv(a, c) for a, c in COMBOS]


def put_header(ws, row: int, names: list[str], start_col: int = 1) -> None:
    for i, name in enumerate(names):
        c = ws.cell(row=row, column=start_col + i, value=name)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = AL_C
        c.border = BORDER


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


wb = openpyxl.Workbook()
wb.properties.title = TITLE

# ---------------- Sheet 1: 测试集汇总 ----------------
ws = wb.active
ws.title = "测试集汇总"
HDR = ["序号", "任务", "工况", "模型", "验证准确率", "测试准确率", "平衡准确率", "宏平均 F1",
       "训练窗数", "验证窗数", "测试窗数", "测试正确数", "测试错误数", "最优轮次", "参数冻结后评估"]
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR))
ws["A1"] = TITLE
ws["A1"].font = F_TITLE
ws["A1"].alignment = AL_L
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HDR))
ws["A2"] = ("数据来源：datas/experiments_seu/checkpoints_bearing_{cnn,gcn}_resplit/bearing/<工况>/seed_42/"
            "（metrics.json + test_predictions.csv）。协议 seu_time_split：70/20/10 时间有序划分，seed 42，"
            "测试集落在时间轴最后 10%，超参数冻结后一次性评估。")
ws["A2"].font = F_SUB
ws["A2"].alignment = AL_L
ws.row_dimensions[2].height = 16

put_header(ws, 4, HDR)
r = 5
for run in RUNS:
    m = run["m"]
    vals = [r - 4, "轴承", COND_CN[run["cond"]], ARCH_CN[run["arch"]],
            m["val_acc"], m["test_acc"], m["test_balanced_acc"], m["test_macro_f1"],
            sum(m["train_class_counts"].values()), sum(m["val_class_counts"].values()),
            sum(m["test_class_counts"].values()), run["correct"], run["n"] - run["correct"],
            m["best_epoch"], "是" if m["test_evaluated_after_freeze"] else "否"]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = F_BODY
        c.border = BORDER
        c.alignment = AL_C if i not in (2, 3, 4) else AL_C
        if i in (5, 6, 7, 8):
            c.number_format = "0.00%"
        elif i in (9, 10, 11, 12, 13, 14):
            c.number_format = "#,##0"
    ws.cell(row=r, column=4).font = F_BOLD
    r += 1

total_n = sum(x["n"] for x in RUNS)
total_ok = sum(x["correct"] for x in RUNS)
lo = min(x["m"]["test_acc"] for x in RUNS)
hi = max(x["m"]["test_acc"] for x in RUNS)
ws.cell(row=r, column=1, value="合计/区间")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws.cell(row=r, column=6, value=f"{lo:.2%} ~ {hi:.2%}")
ws.cell(row=r, column=11, value=total_n)
ws.cell(row=r, column=12, value=total_ok)
ws.cell(row=r, column=13, value=total_n - total_ok)
for i in range(1, len(HDR) + 1):
    c = ws.cell(row=r, column=i)
    c.font = F_BOLD
    c.fill = FILL_ACCENT
    c.border = BORDER
    c.alignment = AL_C
ws.cell(row=r, column=11).number_format = "#,##0"
ws.cell(row=r, column=12).number_format = "#,##0"
ws.cell(row=r, column=13).number_format = "#,##0"

ws.conditional_formatting.add(f"F5:F{r - 1}",
                              CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=XL_OK_BG),
                                         font=Font(name="微软雅黑", size=10, color=XL_OK_FG)))
ws.conditional_formatting.add(f"F5:F{r - 1}",
                              CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor=XL_WARN_BG),
                                         font=Font(name="微软雅黑", size=10, color=XL_WARN_FG)))
set_widths(ws, {"A": 9, "B": 7, "C": 11, "D": 8, "E": 12, "F": 12, "G": 12, "H": 11,
                "I": 10, "J": 10, "K": 10, "L": 12, "M": 12, "N": 11, "O": 16})
ws.freeze_panes = "A5"

# ---------------- Sheet 2: 混淆矩阵 ----------------
ws2 = wb.create_sheet("混淆矩阵")
ws2.merge_cells("A1:G1")
ws2["A1"] = "轴承测试集混淆矩阵（行=真实类别，列=预测类别；每格为该类别下的窗口数）"
ws2["A1"].font = F_TITLE
ws2["A1"].alignment = AL_L
ws2["A2"] = "测试集每类 130 窗、5 类合计 650 窗。矩阵由 test_predictions.csv 逐窗预测重建，与 metrics.json 完全一致。"
ws2["A2"].font = F_SUB

block_row = 4
for idx, run in enumerate(RUNS, start=1):
    cm = run["cm"]
    ws2.merge_cells(start_row=block_row, start_column=1, end_row=block_row, end_column=7)
    ws2.cell(row=block_row, column=1,
             value=f"{idx}. {COND_CN[run['cond']]} · {ARCH_CN[run['arch']]}　"
                   f"测试准确率 {run['correct'] / run['n']:.2%}（{run['n']} 窗，错 {run['n'] - run['correct']}）")
    ws2.cell(row=block_row, column=1).font = F_BLOCK
    hr = block_row + 1
    put_header(ws2, hr, ["真实＼预测"] + [f"{c}（{CLASS_CN[c]}）" for c in CLASSES] + ["行合计（实际）"])
    for i in range(5):
        rr = hr + 1 + i
        c0 = ws2.cell(row=rr, column=1, value=f"{CLASSES[i]}（{CLASS_CN[CLASSES[i]]}）")
        c0.font = F_BOLD
        c0.border = BORDER
        c0.alignment = AL_L
        for j in range(5):
            c = ws2.cell(row=rr, column=2 + j, value=cm[i][j])
            c.font = F_BODY
            c.border = BORDER
            c.alignment = AL_C
            c.number_format = "#,##0"
            if i == j:
                c.fill = FILL_DIAG
                c.font = F_BOLD
        ct = ws2.cell(row=rr, column=7, value=sum(cm[i]))
        ct.font = F_BOLD
        ct.border = BORDER
        ct.alignment = AL_C
        ct.fill = FILL_TOTAL
        ct.number_format = "#,##0"
    rr = hr + 6
    tl = ws2.cell(row=rr, column=1, value="列合计（预测）")
    tl.font = F_BOLD
    tl.border = BORDER
    tl.alignment = AL_L
    tl.fill = FILL_TOTAL
    for j in range(5):
        c = ws2.cell(row=rr, column=2 + j, value=sum(cm[i][j] for i in range(5)))
        c.font = F_BOLD
        c.border = BORDER
        c.alignment = AL_C
        c.fill = FILL_TOTAL
        c.number_format = "#,##0"
    gt = ws2.cell(row=rr, column=7, value=sum(sum(row) for row in cm))
    gt.font = F_BOLD
    gt.border = BORDER
    gt.alignment = AL_C
    gt.fill = FILL_TOTAL
    gt.number_format = "#,##0"
    block_row = rr + 3

set_widths(ws2, {"A": 22, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 16})

# ---------------- Sheet 3: 逐类指标 ----------------
ws3 = wb.create_sheet("逐类指标")
HDR3 = ["序号", "工况", "模型", "类别", "测试样本数", "正确数", "召回率", "精确率", "F1 分数"]
ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR3))
ws3["A1"] = "轴承测试集逐类指标（由混淆矩阵计算：召回=对角/行合计，精确=对角/列合计）"
ws3["A1"].font = F_TITLE
ws3["A1"].alignment = AL_L
put_header(ws3, 3, HDR3)
r3 = 4
k = 0
for run in RUNS:
    cm = run["cm"]
    for i in range(5):
        k += 1
        rowsum = sum(cm[i])
        colsum = sum(cm[j][i] for j in range(5))
        tp = cm[i][i]
        rec = tp / rowsum if rowsum else None
        pre = tp / colsum if colsum else None
        f1 = 2 * rec * pre / (rec + pre) if rec and pre and (rec + pre) else 0.0
        for ci, v in enumerate([k, COND_CN[run["cond"]], ARCH_CN[run["arch"]],
                                f"{CLASSES[i]}（{CLASS_CN[CLASSES[i]]}）", rowsum, tp, rec, pre, f1], start=1):
            c = ws3.cell(row=r3, column=ci, value=v)
            c.font = F_BODY
            c.border = BORDER
            c.alignment = AL_L if ci == 4 else AL_C
            if ci in (7, 8, 9):
                c.number_format = "0.00%"
                c.font = F_BOLD
            elif ci in (5, 6):
                c.number_format = "#,##0"
        r3 += 1
ws3.conditional_formatting.add(f"I4:I{r3 - 1}",
                               CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor=XL_BAD_BG),
                                          font=Font(name="微软雅黑", size=10, bold=True, color=XL_BAD_FG)))
set_widths(ws3, {"A": 7, "B": 11, "C": 8, "D": 18, "E": 12, "F": 10, "G": 11, "H": 11, "I": 11})
ws3.freeze_panes = "A4"

# ---------------- Sheet 4: 测试集逐样本预测 ----------------
ws4 = wb.create_sheet("测试集逐样本预测")
HDR4 = ["序号", "工况", "模型", "样本序号", "真实标签", "预测标签", "是否正确"]
ws4.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR4))
ws4["A1"] = "轴承测试集逐窗预测明细（4 个组合 × 650 窗 = 2,600 行，原样导出）"
ws4["A1"].font = F_TITLE
ws4["A1"].alignment = AL_L
put_header(ws4, 3, HDR4)
r4 = 4
idx = 0
bad_rows = []
for run in RUNS:
    for p in run["preds"]:
        idx += 1
        tl = CLASSES[int(p["true_label"])]
        pl = CLASSES[int(p["pred_label"])]
        ok = tl == pl
        vals = [idx, COND_CN[run["cond"]], ARCH_CN[run["arch"]], int(p["sample_index"]), tl, pl, "是" if ok else "否"]
        for ci, v in enumerate(vals, start=1):
            c = ws4.cell(row=r4, column=ci, value=v)
            c.font = F_BODY
            c.border = BORDER
            c.alignment = AL_C
        if not ok:
            bad_rows.append(r4)
        r4 += 1
if bad_rows:
    for rr in bad_rows:
        for ci in range(1, len(HDR4) + 1):
            cc = ws4.cell(row=rr, column=ci)
            cc.fill = PatternFill("solid", fgColor=XL_BAD_BG)
            cc.font = Font(name="微软雅黑", size=10, bold=True, color=XL_BAD_FG)
set_widths(ws4, {"A": 9, "B": 11, "C": 8, "D": 11, "E": 13, "F": 13, "G": 11})
ws4.freeze_panes = "A4"
ws4.auto_filter.ref = f"A3:G{r4 - 1}"

# ---------------- Sheet 5: 五折交叉验证 ----------------
ws5 = wb.create_sheet("五折交叉验证")
HDR5 = ["序号", "工况", "模型", "折 1", "折 2", "折 3", "折 4", "折 5", "均值", "标准差", "最低折", "最高折"]
ws5.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR5))
ws5["A1"] = "轴承 rolling-origin 五折时间序列交叉验证（test 准确率，每折测试段不同、train 严格早于 val/test）"
ws5["A1"].font = F_TITLE
ws5["A1"].alignment = AL_L
ws5["A2"] = "5 折 test 段分别落在时间轴 [70%,80%]、[75%,85%]、[80%,90%]、[85%,95%]、[90%,100%]。每折 650 个测试窗，5 折合计 3,250 个测试样本。"
ws5["A2"].font = F_SUB
put_header(ws5, 4, HDR5)
r5 = 5
for i, cv in enumerate(CVS, start=1):
    folds = [f["test_acc"] for f in cv["folds_detail"]]
    vals = [i, COND_CN[cv["condition"]], ARCH_CN[cv["arch"]]] + folds + \
           [cv["test_mean"], cv["test_std"], cv["test_min"], cv["test_max"]]
    for ci, v in enumerate(vals, start=1):
        c = ws5.cell(row=r5, column=ci, value=v)
        c.font = F_BODY
        c.border = BORDER
        c.alignment = AL_C
        if ci >= 4:
            c.number_format = "0.00%"
    for ci in (9, 10, 11, 12):
        ws5.cell(row=r5, column=ci).font = F_BOLD
    r5 += 1
allfolds = [f["test_acc"] for cv in CVS for f in cv["folds_detail"]]
ws5.cell(row=r5, column=1, value="全部 20 折合计")
ws5.merge_cells(start_row=r5, start_column=1, end_row=r5, end_column=8)
ws5.cell(row=r5, column=11, value=min(allfolds))
ws5.cell(row=r5, column=12, value=max(allfolds))
for ci in range(1, len(HDR5) + 1):
    c = ws5.cell(row=r5, column=ci)
    c.font = F_BOLD
    c.fill = FILL_ACCENT
    c.border = BORDER
    c.alignment = AL_C
ws5.cell(row=r5, column=11).number_format = "0.00%"
ws5.cell(row=r5, column=12).number_format = "0.00%"
set_widths(ws5, {"A": 14, "B": 11, "C": 8, "D": 10, "E": 10, "F": 10, "G": 10, "H": 10,
                 "I": 11, "J": 11, "K": 11, "L": 11})

# ---------------- Sheet 6: 数据来源与说明 ----------------
ws6 = wb.create_sheet("数据来源与说明")
ws6.merge_cells("A1:B1")
ws6["A1"] = "数据来源与说明"
ws6["A1"].font = F_TITLE
rows = [
    ("导出对象", "《可对比文献对照表》中「本文」轴承一行：99.85% ~ 100%"),
    ("对应实验", "SEU DDS 轴承数据集，5 分类（health/ball/inner/outer/comb），工况 20 Hz-0 V 与 30 Hz-2 V 各自建模，CNN 与 GCN 两条网络，共 4 个组合"),
    ("评估协议", "seu_time_split：70/20/10 时间有序划分（train 916 / val 261 / test 130，每类），seed 42，测试集落在时间轴最后 10%，超参数冻结后一次性评估"),
    ("原始文件①", str(SEU / "checkpoints_bearing_cnn_resplit/bearing/20_0/seed_42")),
    ("原始文件②", str(SEU / "checkpoints_bearing_cnn_resplit/bearing/30_2/seed_42")),
    ("原始文件③", str(SEU / "checkpoints_bearing_gcn_resplit/bearing/20_0/seed_42")),
    ("原始文件④", str(SEU / "checkpoints_bearing_gcn_resplit/bearing/30_2/seed_42")),
    ("原始文件⑤", str(SEU / "cv") + "  （五折交叉验证 bearing_{cnn,gcn}_{20_0,30_2}.json）"),
    ("每个目录内含", "metrics.json = 汇总指标与混淆矩阵；test_predictions.csv = 650 个测试窗的逐样本「真实标签/预测标签」；val_predictions.csv = 1,305 个验证窗逐样本预测"),
    ("准确率定义", "测试准确率 = 正确窗数 / 650；平衡准确率 = 各类召回率的算术平均；宏平均 F1 = 五类 F1 的算术平均"),
    ("区间含义", "99.85% ~ 100% = 四个组合测试准确率的最小值（0.9985，20-0 的 CNN 与 GCN，各 650 窗错 1）到最大值（1.0000，30-2 的 CNN 与 GCN，650 窗全对）"),
    ("校验情况", "本表「逐样本预测」页由 test_predictions.csv 原样导出，重建的混淆矩阵与准确率和 metrics.json 逐格一致（2026-10-01 已核）"),
    ("稳健性补充", "除上述单次留出外，另做 rolling-origin 五折时间序列交叉验证：20 个测试折无一低于 99.23%，全局最低单折 99.23%（详见「五折交叉验证」页）"),
    ("复现命令", "python scripts/experiments/train_seu.py --config config/seu_bearing_cnn.yaml --seeds 42 --evaluate-test --confirm-test（GCN 同理换 config；五折用 scripts/experiments/cv_seu.py）"),
]
rr = 3
for k, v in rows:
    a = ws6.cell(row=rr, column=1, value=k)
    a.font = F_BOLD
    a.alignment = AL_LW
    a.border = BORDER
    a.fill = FILL_ACCENT
    b = ws6.cell(row=rr, column=2, value=v)
    b.font = F_BODY
    b.alignment = AL_LW
    b.border = BORDER
    rr += 1
set_widths(ws6, {"A": 16, "B": 118})

wb.save(str(OUT))
print("saved:", OUT)
print("sheets:", wb.sheetnames)
print("runs:", [(r["arch"], r["cond"], round(r["m"]["test_acc"], 6), r["correct"], r["n"]) for r in RUNS])
print("bad rows:", bad_rows)
