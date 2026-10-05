# -*- coding: utf-8 -*-
"""查 v4 的表注写法（作为样式参考）"""
import re
import docx

d = docx.Document('../导师补充要求_图表与实验汇总_v4.docx')
for p in d.paragraphs:
    t = p.text.strip()
    if re.match(r'^表\s*\d', t) and len(t) < 80:
        x = p._element.xml
        x = re.sub(r'xmlns:\w+="[^"]*"\s*', '', x)
        print('---', t[:60], '---')
        print(x[:700])
        print()
        break

# 列出 v4 所有表注
print('=== v4 表注清单 ===')
for p in d.paragraphs:
    t = p.text.strip()
    if re.match(r'^表\s*\d+[\u3000\s]', t):
        print(' ', t[:90])
