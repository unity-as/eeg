# -*- coding: utf-8 -*-
"""把《可对比文献对照表_给老师_核校版.docx》打成 v2：只改 7 处已核实的错误，保留原格式。"""
import shutil, sys
import docx

SRC = r"C:/Users/lgt11/Desktop/EEG/EEG/doc/references/可对比文献对照表_给老师_核校版.docx"
DST = r"C:/Users/lgt11/Desktop/EEG/EEG/doc/references/可对比文献对照表_给老师_核校版v2.docx"

shutil.copyfile(SRC, DST)
d = docx.Document(DST)

log = []


def set_run(p, idx, old, new, tag):
    r = p.runs[idx]
    assert r.text == old, f"[{tag}] run{idx} 期望 {old!r} 实际 {r.text!r}"
    r.text = new
    log.append(tag)


paras = d.paragraphs

# 1) 数据划分：65:15:10:10（四段）-> 70:20:10 + 留出复核
p = paras[17]
set_run(p, 6, " / ", "", "P17.drop")          # 去掉“测试 / ”后多余的斜杠
set_run(p, 7, "留出四段，比例", "三段，比例", "P17.a")
set_run(p, 8, " 65 : 15 : 10 : 10", " 70 : 20 : 10", "P17.b")
set_run(p, 9, "，测试段与留出段均为密封数据；",
        "，测试段为密封数据（参数冻结后仅评估一次）；另做一组独立的四段留出复核（训练 65% / 验证 15% / 测试 10% / 留出 10%）；", "P17.c")

# 2) 表 0 说明句
p = paras[26]
set_run(p, 0, "下表数值出自本文实验记录，测试集与留出集均为密封数据，随机种子固定为",
        "下表数值出自本文实验记录（70 : 20 : 10 时间有序划分下的密封测试集），随机种子固定为", "P26")

# 3) 表 2 脚注 ＊＊ 补充取值口径
p = paras[32]
set_run(p, 6, "。\n",
        "；表中 91.76 ~ 97.25% 取自该文 SVMD 熵特征在工况 II 下四个分类器的结果。\n", "P32")

# 4) 文献出处 第 8 条（Mathematics 1393 题名）
p = paras[53]
set_run(p, 0,
        "Zhou S, Yu X, Li X, et al. CS-LAMRNet: A Lightweight Attention-Based Multi-Resolution Network for Fault Diagnosis[J]. Mathematics, 2025, 13(9): 1393. DOI: 10.3390/math13091393",
        "Zhou S, Yu X, Li X, Wang Y, Ji K, Ren Z. Gearbox Fault Diagnosis Based on Compressed Sensing and Multi-Scale Residual Network with Lightweight Attention Mechanism[J]. Mathematics, 2025, 13(9): 1393. DOI: 10.3390/math13091393", "P53")

# 5) 文献出处 第 9 条（期刊与卷页）
p = paras[54]
set_run(p, 0,
        "Hierarchical Convolution-Transformer Framework for Gear Fault Diagnosis Under Severe Noise[J]. IEEE Transactions on Instrumentation and Measurement, 2025. DOI: 10.1109/TIM.2025.11037744",
        "He Q, Kang B, Fan S, Li X. Hierarchical Convolution-Transformer Framework for Gear Fault Diagnosis Under Severe Noise[J]. IEEE Access, 2025, 13: 105492-105504. DOI: 10.1109/ACCESS.2025.3580424", "P54")

# 6) 文献出处 第 12 条（ADSRN 题名）
p = paras[57]
set_run(p, 0,
        "ADSRN: Few-Shot Bearing Fault Diagnosis Based on Self-Supervised Learning[J]. Machines, 2025, 13(3): 193. DOI: 10.3390/machines13030193",
        "Jiang Y, Lu M, Dong Z, et al. Adaptive Deeping Siamese Residual Network: A Novel Model for Few-Shot Bearing Fault Diagnosis[J]. Machines, 2025, 13(3): 193. DOI: 10.3390/machines13030193", "P57")

# 7) 表 1 出处栏：IEEE 2025 -> IEEE Access 2025
cell = d.tables[1].rows[16].cells[5]
set_run(cell.paragraphs[0], 0, "IEEE 2025", "IEEE Access 2025", "T1r16c5")

d.save(DST)
print("OK 修改 %d 处:" % len(log), log)
