# -*- coding: utf-8 -*-
"""核对 v3 改动结果"""
import re
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

F = '可对比文献对照表_给老师_核校版v3.docx'
d = docx.Document(F)
print('段落', len(d.paragraphs), '| 表格', len(d.tables))
print('=' * 70)

items = []
for ch in d.element.body.iterchildren():
    if ch.tag.endswith('}p'):
        items.append(('P', Paragraph(ch, d)))
    elif ch.tag.endswith('}tbl'):
        items.append(('T', Table(ch, d)))

print('=== 文档正文顺序（只列非空段与表） ===')
for i, (k, o) in enumerate(items):
    if k == 'P':
        t = o.text.strip()
        if t:
            print(f'{i:3} P {t[:100]}')
    else:
        print(f'{i:3} >>>TABLE {len(o.rows)}x{len(o.columns)}<<<')

print()
print('=== 表注检查 ===')
caps = [o.text.strip() for k, o in items if k == 'P' and re.match(r'^表\s*\d+[\u3000\s]', o.text.strip())]
for c in caps:
    print(' ', c)
print('表注数:', len(caps))

print()
print('=== 引用编号检查 ===')
full = '\n'.join(o.text for k, o in items if k == 'P')
for m in re.finditer(r'.{10}表\s*\d+.{14}', full):
    print('  ', repr(m.group(0)))
